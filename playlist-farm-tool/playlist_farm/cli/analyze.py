#!/usr/bin/env python3
"""
Analysis tool for playlist database
"""

from openpyxl import load_workbook
import os
import sys
from playlist_farm.core.config import DEFAULT_EXCEL_FILE


def main():
    """Analyze the playlist database"""
    print("""
┌────────────────────────────────────────────────────┐
│ ░█▀█░█░░░█▀█░█░█░█░░░▀█▀░█▀▀░▀█▀░░░█▀▀░█▀█░█▀▄░█▄█ │
│ ░█▀▀░█░░░█▀█░░█░░█░░░░█░░▀▀█░░█░░░░█▀▀░█▀█░█▀▄░█░█ │
│ ░▀░░░▀▀▀░▀░▀░░▀░░▀▀▀░▀▀▀░▀▀▀░░▀░░░░▀░░░▀░▀░▀░▀░▀░▀ │
└────────────────────────────────────────────────────┘
""")

    excel_file = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_EXCEL_FILE

    if not os.path.exists(excel_file):
        print(f"❌ File not found: {excel_file}")
        return

    print(f"\n📊 Analyzing: {excel_file}\n")
    print("=" * 70)

    wb = load_workbook(excel_file)
    ws = wb.active

    # Statistics
    total = ws.max_row - 1  # Exclude header
    with_email = 0
    with_instagram = 0
    with_twitter = 0
    with_any_contact = 0

    total_followers = 0
    playlists_with_contact = []

    # Analyze rows
    for row in range(2, ws.max_row + 1):
        email = ws.cell(row, 10).value
        instagram = ws.cell(row, 11).value
        twitter = ws.cell(row, 12).value
        followers = ws.cell(row, 7).value or 0

        total_followers += followers

        if email:
            with_email += 1
        if instagram:
            with_instagram += 1
        if twitter:
            with_twitter += 1
        if email or instagram or twitter:
            with_any_contact += 1
            playlists_with_contact.append({
                'name': ws.cell(row, 1).value,
                'curator': ws.cell(row, 4).value,
                'followers': followers,
                'email': email,
                'instagram': instagram,
                'twitter': twitter,
                'url': ws.cell(row, 3).value
            })

    # Print statistics
    print(f"\n📈 DATABASE STATISTICS\n")
    print(f"  Total Playlists:         {total:,}")
    print(f"  Total Followers:         {total_followers:,}")
    print(f"  Average Followers:       {total_followers//total if total > 0 else 0:,}")
    print(f"\n📧 CONTACT INFORMATION\n")
    print(f"  With Email:             {with_email:,} ({with_email/total*100 if total > 0 else 0:.1f}%)")
    print(f"  With Instagram:         {with_instagram:,} ({with_instagram/total*100 if total > 0 else 0:.1f}%)")
    print(f"  With Twitter:           {with_twitter:,} ({with_twitter/total*100 if total > 0 else 0:.1f}%)")
    print(f"  With ANY Contact:       {with_any_contact:,} ({with_any_contact/total*100 if total > 0 else 0:.1f}%)")

    # Show top playlists with contact
    if playlists_with_contact:
        print(f"\n🎯 TOP PLAYLISTS WITH CONTACT INFO (sorted by followers)\n")
        print("=" * 70)

        sorted_playlists = sorted(playlists_with_contact, key=lambda x: x['followers'], reverse=True)

        for i, p in enumerate(sorted_playlists[:20], 1):
            print(f"\n{i}. {p['name'][:50]}")
            print(f"   Curator: {p['curator']}")
            print(f"   Followers: {p['followers']:,}")

            if p['email']:
                print(f"   📧 {p['email']}")
            if p['instagram']:
                print(f"   📱 @{p['instagram']}")
            if p['twitter']:
                print(f"   🐦 @{p['twitter']}")

    print("\n" + "=" * 70)
    print(f"\n✅ Analysis complete!\n")


if __name__ == "__main__":
    main()
