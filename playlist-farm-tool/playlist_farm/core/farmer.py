"""Core farming logic for Playlist Farm"""

import requests
import base64
import re
import time
import os
from datetime import datetime
from typing import List, Dict, Set, Optional, Callable
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from playlist_farm.core.config import CLIENT_ID, CLIENT_SECRET, DEFAULT_EXCEL_FILE

try:
    from playlist_farm.core.db_writer import DatabaseWriter
    DB_SUPPORT = True
except ImportError:
    DB_SUPPORT = False


class PlaylistFarmer:
    """Main class for farming Spotify playlist curator information"""

    def __init__(self, excel_file: str = DEFAULT_EXCEL_FILE, database_url: Optional[str] = None, log_callback: Optional[Callable] = None):
        self.access_token = None
        self.base_url = "https://api.spotify.com/v1"
        self.excel_file = excel_file
        self.existing_ids: Set[str] = set()
        self.workbook = None
        self.worksheet = None
        self.log_callback = log_callback or print

        # Database support
        self.database_url = database_url
        self.db_writer: Optional['DatabaseWriter'] = None
        self.use_database = False

        # Statistics
        self.stats = {
            'total_processed': 0,
            'new_added': 0,
            'duplicates_skipped': 0,
            'with_contact': 0,
            'with_email': 0,
            'with_instagram': 0,
            'with_twitter': 0,
            'total_followers': 0,
            'errors': 0
        }

    def log(self, message: str):
        """Log a message"""
        if self.log_callback:
            self.log_callback(message)

    def authenticate(self):
        """Get Spotify access token"""
        auth_url = "https://accounts.spotify.com/api/token"
        auth_header = base64.b64encode(
            f"{CLIENT_ID}:{CLIENT_SECRET}".encode()
        ).decode()

        headers = {
            "Authorization": f"Basic {auth_header}",
            "Content-Type": "application/x-www-form-urlencoded"
        }

        response = requests.post(auth_url, headers=headers, data={"grant_type": "client_credentials"})
        response.raise_for_status()
        self.access_token = response.json()["access_token"]

    def init_database(self):
        """Initialize database connection if database_url is provided"""
        if not self.database_url:
            return False

        if not DB_SUPPORT:
            self.log("⚠ Database support not available (missing dependencies)")
            return False

        try:
            self.db_writer = DatabaseWriter(self.database_url, log_callback=self.log)
            if self.db_writer.connect():
                self.use_database = True
                # Merge existing IDs from database
                self.existing_ids.update(self.db_writer.existing_ids)
                return True
        except Exception as e:
            self.log(f"⚠ Database initialization failed: {e}")
            self.db_writer = None
            return False

    def init_excel(self):
        """Initialize or load existing Excel file"""
        if os.path.exists(self.excel_file):
            self.log(f"📂 Loading existing Excel database: {self.excel_file}")
            self.workbook = load_workbook(self.excel_file)
            self.worksheet = self.workbook.active

            # Load existing playlist IDs to avoid duplicates
            for row in range(2, self.worksheet.max_row + 1):
                playlist_id = self.worksheet.cell(row, 2).value
                if playlist_id:
                    self.existing_ids.add(playlist_id)

            self.log(f"✓ Loaded {len(self.existing_ids):,} existing playlists from Excel")
        else:
            self.log(f"📝 Creating new Excel database: {self.excel_file}")
            self.workbook = Workbook()
            self.worksheet = self.workbook.active
            self.worksheet.title = "Playlists"

            # Create headers
            headers = [
                "Playlist Name", "Playlist ID", "Playlist URL",
                "Curator Name", "Curator ID", "Curator URL",
                "Followers", "Total Tracks", "Is Public",
                "Email", "Instagram", "Twitter",
                "Description", "Genre Searched", "Date Added"
            ]

            for col, header in enumerate(headers, 1):
                cell = self.worksheet.cell(1, col, header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")

            # Set column widths
            widths = {
                'A': 40, 'B': 25, 'C': 50, 'D': 30, 'E': 25, 'F': 50,
                'G': 12, 'H': 12, 'I': 10, 'J': 35, 'K': 25, 'L': 25,
                'M': 60, 'N': 20, 'O': 20
            }
            for col, width in widths.items():
                self.worksheet.column_dimensions[col].width = width

            self.save_excel()

        # Initialize database if configured
        if self.database_url:
            self.init_database()

    def save_excel(self):
        """Save Excel file"""
        self.workbook.save(self.excel_file)

    def search_playlists(self, query: str, limit: int = 50) -> List[Dict]:
        """Search for playlists"""
        if not self.access_token:
            self.authenticate()

        headers = {"Authorization": f"Bearer {self.access_token}"}
        all_playlists = []
        offset = 0

        while len(all_playlists) < limit:
            batch_size = min(50, limit - len(all_playlists))

            try:
                response = requests.get(
                    f"{self.base_url}/search",
                    headers=headers,
                    params={
                        "q": query,
                        "type": "playlist",
                        "limit": batch_size,
                        "offset": offset
                    }
                )
                response.raise_for_status()

                batch = response.json()["playlists"]["items"]
                if not batch:
                    break

                all_playlists.extend([p for p in batch if p])
                offset += batch_size

                if len(batch) < batch_size:
                    break

            except Exception as e:
                self.log(f"⚠ Error searching: {e}")
                break

        return all_playlists

    def get_playlist_details(self, playlist_id: str) -> Dict:
        """Get detailed playlist information"""
        if not self.access_token:
            self.authenticate()

        headers = {"Authorization": f"Bearer {self.access_token}"}
        response = requests.get(f"{self.base_url}/playlists/{playlist_id}", headers=headers)
        response.raise_for_status()
        return response.json()

    def extract_contact_info(self, description: str) -> Dict[str, str]:
        """Extract contact information from description"""
        if not description:
            return {}

        contacts = {}

        # Email
        email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
        emails = re.findall(email_pattern, description)
        if emails:
            contacts['email'] = emails[0]

        # Instagram
        ig_pattern = r'(?:@|instagram\.com/|ig:?\s*)([A-Za-z0-9._]{1,30})'
        ig_matches = re.findall(ig_pattern, description.lower())
        if ig_matches:
            contacts['instagram'] = ig_matches[0]

        # Twitter/X
        twitter_pattern = r'(?:twitter\.com/|x\.com/)([A-Za-z0-9._]{1,15})'
        twitter_matches = re.findall(twitter_pattern, description.lower())
        if twitter_matches:
            contacts['twitter'] = twitter_matches[0]

        return contacts

    def add_playlist(self, playlist_data: Dict, genre: str):
        """Add a single playlist to Excel and database (if configured)"""
        next_row = self.worksheet.max_row + 1

        # Add data
        cells = [
            playlist_data.get('playlist_name', ''),
            playlist_data.get('playlist_id', ''),
            playlist_data.get('playlist_url', ''),
            playlist_data.get('curator_name', ''),
            playlist_data.get('curator_id', ''),
            playlist_data.get('curator_url', ''),
            playlist_data.get('followers', 0),
            playlist_data.get('total_tracks', 0),
            'Yes' if playlist_data.get('is_public') else 'No',
            playlist_data.get('email', ''),
            playlist_data.get('instagram', ''),
            playlist_data.get('twitter', ''),
            playlist_data.get('description', ''),
            genre,
            datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        ]

        for col, value in enumerate(cells, 1):
            self.worksheet.cell(next_row, col, value)

        # Highlight rows with contact info
        if playlist_data.get('email') or playlist_data.get('instagram'):
            for col in range(1, 16):
                self.worksheet.cell(next_row, col).fill = PatternFill(
                    start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"
                )

        self.save_excel()

        # Also add to database if enabled
        if self.use_database and self.db_writer:
            try:
                self.db_writer.add_playlist(playlist_data, genre)
            except Exception as e:
                self.log(f"⚠ Database write error: {e}")

    def process_playlist(self, playlist: Dict, genre: str, callback: Optional[Callable] = None) -> bool:
        """Process a single playlist and add to database"""
        playlist_id = playlist.get('id')

        # Skip if already processed
        if playlist_id in self.existing_ids:
            self.stats['duplicates_skipped'] += 1
            self.stats['total_processed'] += 1
            return False

        try:
            # Get details
            details = self.get_playlist_details(playlist_id)

            playlist_data = {
                'playlist_name': details.get('name', 'Unknown'),
                'playlist_id': details.get('id', ''),
                'playlist_url': details.get('external_urls', {}).get('spotify', ''),
                'curator_name': details.get('owner', {}).get('display_name', 'Unknown'),
                'curator_id': details.get('owner', {}).get('id', ''),
                'curator_url': details.get('owner', {}).get('external_urls', {}).get('spotify', ''),
                'followers': details.get('followers', {}).get('total', 0),
                'total_tracks': details.get('tracks', {}).get('total', 0),
                'description': details.get('description', ''),
                'is_public': details.get('public', False),
            }

            # Extract contacts
            contacts = self.extract_contact_info(playlist_data['description'])
            playlist_data.update(contacts)

            # Add to data sources (Excel + Database if configured)
            self.add_playlist(playlist_data, genre)

            # Update tracking
            self.existing_ids.add(playlist_id)
            self.stats['new_added'] += 1
            self.stats['total_processed'] += 1
            self.stats['total_followers'] += playlist_data['followers']

            if contacts:
                self.stats['with_contact'] += 1
                if contacts.get('email'):
                    self.stats['with_email'] += 1
                if contacts.get('instagram'):
                    self.stats['with_instagram'] += 1
                if contacts.get('twitter'):
                    self.stats['with_twitter'] += 1

            # Print progress
            has_contact = "📧" if contacts else "  "
            self.log(f"{has_contact} Added: {playlist_data['playlist_name'][:50]} | {playlist_data['curator_name']} | {playlist_data['followers']:,} followers")

            if callback:
                callback(playlist_data, contacts)

            return True

        except Exception as e:
            self.stats['errors'] += 1
            self.stats['total_processed'] += 1
            self.log(f"⚠ Error processing playlist: {e}")
            return False

    def farm_genre(self, genre: str, limit_per_genre: int = 50, progress_callback: Optional[Callable] = None):
        """Farm playlists for a specific genre"""
        self.log(f"\n🌾 Farming genre: '{genre}'")
        self.log(f"{'=' * 60}")

        playlists = self.search_playlists(genre, limit=limit_per_genre)

        if not playlists:
            self.log(f"  No playlists found for '{genre}'")
            return

        self.log(f"  Found {len(playlists)} playlists, processing...")

        for idx, playlist in enumerate(playlists, 1):
            self.process_playlist(playlist, genre)

            if progress_callback:
                progress_callback(idx, len(playlists))

            # Rate limiting
            time.sleep(0.2)

    def get_database_stats(self) -> Dict:
        """Get database statistics"""
        if not self.worksheet:
            return {'total': 0, 'with_contact': 0, 'contact_pct': 0, 'total_followers': 0}

        total = self.worksheet.max_row - 1
        with_contact = 0
        total_followers = 0

        for row in range(2, self.worksheet.max_row + 1):
            email = self.worksheet.cell(row, 10).value
            instagram = self.worksheet.cell(row, 11).value
            followers = self.worksheet.cell(row, 7).value or 0

            total_followers += followers

            if email or instagram:
                with_contact += 1

        return {
            'total': total,
            'with_contact': with_contact,
            'contact_pct': (with_contact / total * 100) if total > 0 else 0,
            'total_followers': total_followers
        }
